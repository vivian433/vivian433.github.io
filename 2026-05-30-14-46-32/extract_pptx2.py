
import sys
import os

# 设置输出编码
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation

def extract_pptx(path):
    try:
        prs = Presentation(path)
        texts = []
        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t and len(t) > 1:
                            slide_texts.append(t)
            if slide_texts:
                texts.append(f"[第{i+1}页] " + " | ".join(slide_texts))
        return "\n".join(texts)
    except Exception as e:
        return f"ERROR: {e}"

files = [
    ("项目1-认识安装linux", "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目1 认识与安装linux操作系统.pptx"),
    ("项目2-基础命令", "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目2 linux基础命令操作.pptx"),
    ("项目3-用户组管理", "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目3 管理Linux服务器的用户和组.pptx"),
    ("项目4-文件系统", "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目4 配置与管理文件系统.pptx"),
    ("项目6-软件包管理", "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目6 软件包管理.pptx"),
    ("第6章-网络配置ssh", "E:/1工作/1课程/1linux操作系统/2025-2026-2/第6章 网络配置和使用ssh服务.pptx"),
]

out_dir = "E:/1工作/workbuddy/2026-05-30-14-46-32/pptx_texts"
os.makedirs(out_dir, exist_ok=True)

for name, path in files:
    content = extract_pptx(path)
    out_path = f"{out_dir}/{name}.txt"
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(content)
    print(f"已提取: {name} -> {len(content)} 字符")
