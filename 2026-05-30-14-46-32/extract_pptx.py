
import sys
from pptx import Presentation
from pptx.util import Inches
import os

def extract_pptx(path):
    try:
        prs = Presentation(path)
        texts = []
        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            slide_texts.append(t)
            if slide_texts:
                texts.append(f"[第{i+1}页] " + " | ".join(slide_texts))
        return "\n".join(texts)
    except Exception as e:
        return f"ERROR: {e}"

files = [
    "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目1 认识与安装linux操作系统.pptx",
    "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目2 linux基础命令操作.pptx",
    "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目3 管理Linux服务器的用户和组.pptx",
    "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目4 配置与管理文件系统.pptx",
    "E:/1工作/1课程/1linux操作系统/2025-2026-2/项目6 软件包管理.pptx",
    "E:/1工作/1课程/1linux操作系统/2025-2026-2/第6章 网络配置和使用ssh服务.pptx",
]

for f in files:
    print(f"\n{'='*60}")
    print(f"文件: {os.path.basename(f)}")
    print('='*60)
    content = extract_pptx(f)
    # 限制输出长度
    if len(content) > 8000:
        content = content[:8000] + "\n...(截断)..."
    print(content)
